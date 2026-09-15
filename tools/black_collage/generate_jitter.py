#!/usr/bin/env python3
"""Generate sticker jitter GIFs + layered animated PPTX (no zip).

Paths are relative to this toolkit folder. Requires cutouts in assets/
(run generate_slides.py first, or copy cutout_*.png into assets/).
"""

from __future__ import annotations

import math
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from lxml import etree

from generate_slides import (
    ASSETS,
    ROOT,
    W,
    H,
    PINK,
    FONT_TITLE,
    FONT_LABEL,
    load_font,
    make_sticker,
    rotate_keep,
    draw_hand_text,
    draw_arrow,
    draw_heart,
    draw_star,
    draw_spark,
    new_canvas,
    flatten,
)

FPS = 14
DURATION = 2.4
N_FRAMES = int(FPS * DURATION)


def wobble(t: float, phase: float, amp_deg: float = 2.2, amp_px: float = 4.0):
    rot = amp_deg * math.sin(2 * math.pi * t + phase) + 0.45 * amp_deg * math.sin(
        4 * math.pi * t + phase * 1.3
    )
    dx = amp_px * math.sin(2 * math.pi * t + phase + 0.7) + 0.35 * amp_px * math.cos(
        6 * math.pi * t + phase
    )
    dy = amp_px * math.cos(2 * math.pi * t + phase * 0.9) + 0.3 * amp_px * math.sin(
        4 * math.pi * t + 1.1
    )
    return rot, dx, dy


def soft_wobble(t: float, phase: float):
    return wobble(t, phase, amp_deg=1.0, amp_px=2.0)


def paste_free(canvas: Image.Image, sticker: Image.Image, xy: tuple[float, float], anchor="lt"):
    x, y = xy
    w, h = sticker.size
    if anchor == "center":
        x -= w // 2
        y -= h // 2
    else:
        if "c" in anchor and "l" not in anchor and "r" not in anchor:
            x -= w // 2
        if anchor.endswith("r") or anchor.startswith("r"):
            x -= w
        if "b" in anchor:
            y -= h
    x = int(round(x))
    y = int(round(y))
    x = max(-w // 4, min(x, W - w + w // 4))
    y = max(-h // 4, min(y, H - h + h // 4))
    cx0, cy0 = max(0, x), max(0, y)
    sx0, sy0 = cx0 - x, cy0 - y
    cx1, cy1 = min(W, x + w), min(H, y + h)
    if cx1 <= cx0 or cy1 <= cy0:
        return (x, y, x + w, y + h)
    crop = sticker.crop((sx0, sy0, sx0 + (cx1 - cx0), sy0 + (cy1 - cy0)))
    canvas.alpha_composite(crop, (cx0, cy0))
    return (x, y, x + w, y + h)


def prepare_stickers():
    for key in ("child", "cake", "artist", "dress", "chili"):
        p = ASSETS / f"cutout_{key}.png"
        if not p.exists():
            raise FileNotFoundError(f"Missing {p}. Run generate_slides.py first or copy cutouts.")

    child = Image.open(ASSETS / "cutout_child.png").convert("RGBA")
    cake = Image.open(ASSETS / "cutout_cake.png").convert("RGBA")
    artist = Image.open(ASSETS / "cutout_artist.png").convert("RGBA")
    dress = Image.open(ASSETS / "cutout_dress.png").convert("RGBA")
    chili = Image.open(ASSETS / "cutout_chili.png").convert("RGBA")

    st = {}
    st["child"] = make_sticker(child, border=20, jitter_amp=6.2, seed=21, max_h=880)
    st["cake"] = make_sticker(cake, border=16, jitter_amp=5.4, seed=3, max_h=500, max_w=640)
    st["artist"] = make_sticker(artist, border=15, jitter_amp=5.1, seed=8, max_h=390, max_w=560)
    st["dress"] = make_sticker(dress, border=17, jitter_amp=5.8, seed=12, max_h=740)
    st["dress3"] = make_sticker(dress, border=18, jitter_amp=5.7, seed=19, max_h=900)
    st["chili"] = make_sticker(chili, border=16, jitter_amp=5.4, seed=22, max_h=540, max_w=760)
    return st


def build_slide1_frame(st, t: float) -> Image.Image:
    s = new_canvas()
    r, dx, dy = wobble(t, 0.25, amp_deg=2.5, amp_px=5.5)
    child = rotate_keep(st["child"], -5.0 + r)
    box_c = paste_free(s, child, (620 + dx, 545 + dy), anchor="center")

    f_title = load_font(FONT_TITLE, 120)
    f_label = load_font(FONT_LABEL, 50)
    f_year = load_font(FONT_TITLE, 58)
    tr, tdx, tdy = soft_wobble(t, 1.4)
    draw_hand_text(s, "小时候～", (1180 + tdx, 40 + tdy), f_title, angle=-3.5 + tr * 0.4, jitter=2.5, seed=4)
    draw_hand_text(s, "～", (1580 + tdx * 0.5, 210 + tdy * 0.5), f_year, fill=PINK, angle=14 + tr * 0.3, jitter=1.2, seed=9)
    lr, ldx, ldy = soft_wobble(t, 2.8)
    draw_hand_text(s, "小小的你", (1320 + ldx, 700 + ldy), f_label, angle=7.0 + lr * 0.4, jitter=1.6, seed=6)
    draw_arrow(
        s,
        (1310 + ldx, 750 + ldy),
        (box_c[2] - 40, (box_c[1] + box_c[3]) // 2 + 30),
        dotted=True,
        width=3,
        seed=2,
    )
    hr, hdx, hdy = soft_wobble(t, 0.9)
    draw_heart(s, int(1160 + hdx), int(300 + hdy), s=17, angle=-16 + hr)
    draw_star(s, int(1700 - hdx), int(880 + hdy), s=14, angle=22 + hr)
    draw_spark(s, int(220 + hdx * 0.5), int(160 - hdy * 0.5), s=12)
    draw_heart(s, int(1640 + hdy), int(500 + hdx), s=13, angle=18)
    draw_spark(s, 1080, 980, s=9)
    return flatten(s)


def build_slide2_frame(st, t: float) -> Image.Image:
    s = new_canvas()
    r1, dx1, dy1 = wobble(t, 0.1, amp_deg=2.3, amp_px=5.0)
    r2, dx2, dy2 = wobble(t, 2.1, amp_deg=2.6, amp_px=5.5)
    r3, dx3, dy3 = wobble(t, 4.0, amp_deg=2.2, amp_px=4.8)

    cake = rotate_keep(st["cake"], -7.5 + r1)
    art = rotate_keep(st["artist"], 6.8 + r2)
    dress = rotate_keep(st["dress"], 3.5 + r3)

    box_cake = paste_free(s, cake, (40 + dx1, 300 + dy1), anchor="lt")
    box_art = paste_free(s, art, (680 + dx2, 95 + dy2), anchor="lt")
    box_dress = paste_free(s, dress, (1310 + dx3, 170 + dy3), anchor="lt")

    f_title2 = load_font(FONT_TITLE, 112)
    f_lab = load_font(FONT_LABEL, 46)
    tr, tdx, tdy = soft_wobble(t, 1.1)
    draw_hand_text(s, "长大啦！！", (36 + tdx, 28 + tdy), f_title2, angle=-5.8 + tr * 0.35, jitter=2.3, seed=11)

    lr1, lx1, ly1 = soft_wobble(t, 3.0)
    draw_hand_text(s, "吃蛋糕", (80 + lx1, 920 + ly1), f_lab, angle=-6 + lr1 * 0.3, jitter=1.4, seed=13)
    draw_arrow(s, (220 + lx1, 910 + ly1), ((box_cake[0] + box_cake[2]) // 2, box_cake[3] - 24), dotted=True, seed=4)

    lr2, lx2, ly2 = soft_wobble(t, 3.5)
    draw_hand_text(s, "画画", (1040 + lx2, 70 + ly2), f_lab, angle=8 + lr2 * 0.3, jitter=1.3, seed=14)
    draw_arrow(s, (1100 + lx2, 145 + ly2), (box_art[2] - 40, box_art[1] + 70), dotted=True, seed=5)

    lr3, lx3, ly3 = soft_wobble(t, 4.2)
    draw_hand_text(s, "穿搭", (1688 + lx3, 90 + ly3), f_lab, angle=8 + lr3 * 0.3, jitter=1.4, seed=15)
    draw_arrow(s, (1740 + lx3, 165 + ly3), (box_dress[0] + 90, box_dress[1] + 50), dotted=True, seed=6)

    hr, hdx, hdy = soft_wobble(t, 0.6)
    draw_heart(s, int(620 + hdx), int(180 + hdy), s=16, angle=-22 + hr)
    draw_heart(s, int(1200 - hdx), int(980 + hdy), s=13, angle=14)
    draw_star(s, int(1120 + hdy), int(540 + hdx), s=13, angle=25 + hr)
    draw_spark(s, int(640 + hdx * 0.4), 980, s=11)
    draw_star(s, 1820, 980, s=12, angle=-12)
    draw_spark(s, 980, 980, s=9)
    return flatten(s)


def build_slide3_frame(st, t: float) -> Image.Image:
    s = new_canvas()
    r1, dx1, dy1 = wobble(t, 0.8, amp_deg=2.4, amp_px=5.2)
    r2, dx2, dy2 = wobble(t, 3.3, amp_deg=2.5, amp_px=5.0)

    chili = rotate_keep(st["chili"], 6.2 + r1)
    dress = rotate_keep(st["dress3"], -5.8 + r2)

    box_ch = paste_free(s, chili, (70 + dx1, 200 + dy1), anchor="lt")
    box_dr = paste_free(s, dress, (1200 + dx2, 80 + dy2), anchor="lt")

    f_title = load_font(FONT_TITLE, 120)
    f_lab3 = load_font(FONT_LABEL, 44)
    tr, tdx, tdy = soft_wobble(t, 1.7)
    draw_hand_text(s, "现在的你", (700 + tdx, 36 + tdy), f_title, angle=-2.6 + tr * 0.35, jitter=2.3, seed=17)
    draw_hand_text(
        s,
        "认真生活～闪闪发光",
        (60 + tdx * 0.6, 910 + tdy * 0.6),
        load_font(FONT_TITLE, 64),
        angle=-3.2 + tr * 0.25,
        jitter=1.6,
        seed=18,
    )
    lr, lx, ly = soft_wobble(t, 2.4)
    draw_hand_text(s, "闪光日常", (760 + lx, 250 + ly), f_lab3, angle=10 + lr * 0.3, jitter=1.3, seed=20)
    draw_arrow(s, (860 + lx, 320 + ly), (box_ch[2] - 30, box_ch[1] + 70), dotted=True, seed=8)

    hr, hdx, hdy = soft_wobble(t, 0.4)
    draw_heart(s, int(1080 + hdx), int(150 + hdy), s=18, angle=-12 + hr)
    draw_heart(s, int(520 - hdx), int(140 + hdy), s=13, angle=22)
    draw_star(s, int(1740 + hdy), int(70 + hdx), s=15, angle=18 + hr)
    draw_spark(s, int(980 + hdx * 0.5), int(620 + hdy * 0.5), s=13)
    draw_star(s, 1780, 980, s=13, angle=-8)
    draw_heart(s, 1100, 980, s=12, angle=16)
    return flatten(s)


def save_gif(frames: list[Image.Image], path: Path, fps: int = FPS):
    preview = []
    for im in frames:
        small = im.resize((960, 540), Image.Resampling.LANCZOS)
        preview.append(small.convert("P", palette=Image.ADAPTIVE, colors=128))
    duration_ms = int(1000 / fps)
    preview[0].save(
        path,
        save_all=True,
        append_images=preview[1:],
        duration=duration_ms,
        loop=0,
        optimize=True,
        disposal=2,
    )
    print(f"wrote {path} ({path.stat().st_size / 1e6:.2f} MB)")


def save_mp4(frames: list[Image.Image], path: Path, fps: int = FPS):
    import imageio.v2 as imageio

    arrs = [np.array(im.convert("RGB")) for im in frames]
    writer = imageio.get_writer(
        str(path), fps=fps, codec="libx264", quality=7, pixelformat="yuv420p", macro_block_size=1
    )
    for a in arrs:
        writer.append_data(a)
    writer.close()
    print(f"wrote {path} ({path.stat().st_size / 1e6:.2f} MB)")


NSMAP = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}


def inject_teeter_timing(slide, shape_ids: list[int]):
    sld = slide._element
    for old in sld.findall(f"{{{NSMAP['p']}}}timing"):
        sld.remove(old)

    p = NSMAP["p"]
    timing = etree.SubElement(sld, f"{{{p}}}timing")
    tn_lst = etree.SubElement(timing, f"{{{p}}}tnLst")
    par0 = etree.SubElement(tn_lst, f"{{{p}}}par")
    ctn0 = etree.SubElement(
        par0, f"{{{p}}}cTn", {"id": "1", "dur": "indefinite", "restart": "never", "nodeType": "tmRoot"}
    )
    child_tn = etree.SubElement(ctn0, f"{{{p}}}childTnLst")
    seq = etree.SubElement(child_tn, f"{{{p}}}seq", {"concurrent": "1", "nextAc": "seek"})
    ctn_seq = etree.SubElement(
        seq, f"{{{p}}}cTn", {"id": "2", "dur": "indefinite", "nodeType": "mainSeq"}
    )
    st_cond = etree.SubElement(ctn_seq, f"{{{p}}}stCondLst")
    etree.SubElement(st_cond, f"{{{p}}}cond", {"delay": "0"})
    seq_children = etree.SubElement(ctn_seq, f"{{{p}}}childTnLst")

    next_id = 3
    for spid in shape_ids:
        par = etree.SubElement(seq_children, f"{{{p}}}par")
        ctn = etree.SubElement(
            par, f"{{{p}}}cTn", {"id": str(next_id), "fill": "hold", "nodeType": "withEffect"}
        )
        next_id += 1
        st = etree.SubElement(ctn, f"{{{p}}}stCondLst")
        etree.SubElement(st, f"{{{p}}}cond", {"delay": "0"})
        ch = etree.SubElement(ctn, f"{{{p}}}childTnLst")

        anim_par = etree.SubElement(ch, f"{{{p}}}par")
        anim_ctn = etree.SubElement(
            anim_par,
            f"{{{p}}}cTn",
            {
                "id": str(next_id),
                "presetID": "26",
                "presetClass": "emph",
                "presetSubtype": "0",
                "repeatCount": "indefinite",
                "fill": "hold",
                "nodeType": "withEffect",
            },
        )
        next_id += 1
        st2 = etree.SubElement(anim_ctn, f"{{{p}}}stCondLst")
        etree.SubElement(st2, f"{{{p}}}cond", {"delay": "0"})
        ch2 = etree.SubElement(anim_ctn, f"{{{p}}}childTnLst")

        anim = etree.SubElement(ch2, f"{{{p}}}animEffect", {"transition": "in", "filter": "teeter"})
        c_bhvr = etree.SubElement(anim, f"{{{p}}}cBhvr")
        etree.SubElement(c_bhvr, f"{{{p}}}cTn", {"id": str(next_id), "dur": "2000"})
        next_id += 1
        tgt = etree.SubElement(c_bhvr, f"{{{p}}}tgtEl")
        etree.SubElement(tgt, f"{{{p}}}spTgt", {"spid": str(spid)})

        anim_rot = etree.SubElement(ch2, f"{{{p}}}animRot", {"by": "200000"})
        c_bhvr2 = etree.SubElement(anim_rot, f"{{{p}}}cBhvr")
        etree.SubElement(
            c_bhvr2,
            f"{{{p}}}cTn",
            {"id": str(next_id), "dur": "1000", "autoRev": "1", "repeatCount": "indefinite"},
        )
        next_id += 1
        attr = etree.SubElement(c_bhvr2, f"{{{p}}}attrNameLst")
        an = etree.SubElement(attr, f"{{{p}}}attrName")
        an.text = "r"
        tgt2 = etree.SubElement(c_bhvr2, f"{{{p}}}tgtEl")
        etree.SubElement(tgt2, f"{{{p}}}spTgt", {"spid": str(spid)})


def draw_decor_s1(s1, box_c):
    f_title = load_font(FONT_TITLE, 120)
    f_label = load_font(FONT_LABEL, 50)
    f_year = load_font(FONT_TITLE, 58)
    draw_hand_text(s1, "小时候～", (1180, 40), f_title, angle=-3.5, jitter=2.5, seed=4)
    draw_hand_text(s1, "～", (1580, 210), f_year, fill=PINK, angle=14, jitter=1.2, seed=9)
    draw_hand_text(s1, "小小的你", (1320, 700), f_label, angle=7.0, jitter=1.6, seed=6)
    draw_arrow(s1, (1310, 750), (box_c[2] - 40, (box_c[1] + box_c[3]) // 2 + 30), dotted=True, width=3, seed=2)
    draw_heart(s1, 1160, 300, s=17, angle=-16)
    draw_star(s1, 1700, 880, s=14, angle=22)
    draw_spark(s1, 220, 160, s=12)
    draw_heart(s1, 1640, 500, s=13, angle=18)
    draw_spark(s1, 1080, 980, s=9)


def draw_decor_s2(s2, box_cake, box_art, box_dress):
    f_title2 = load_font(FONT_TITLE, 112)
    f_lab = load_font(FONT_LABEL, 46)
    draw_hand_text(s2, "长大啦！！", (36, 28), f_title2, angle=-5.8, jitter=2.3, seed=11)
    draw_hand_text(s2, "吃蛋糕", (80, 920), f_lab, angle=-6, jitter=1.4, seed=13)
    draw_arrow(s2, (220, 910), ((box_cake[0] + box_cake[2]) // 2, box_cake[3] - 24), dotted=True, seed=4)
    draw_hand_text(s2, "画画", (1040, 70), f_lab, angle=8, jitter=1.3, seed=14)
    draw_arrow(s2, (1100, 145), (box_art[2] - 40, box_art[1] + 70), dotted=True, seed=5)
    draw_hand_text(s2, "穿搭", (1688, 90), f_lab, angle=8, jitter=1.4, seed=15)
    draw_arrow(s2, (1740, 165), (box_dress[0] + 90, box_dress[1] + 50), dotted=True, seed=6)
    draw_heart(s2, 620, 180, s=16, angle=-22)
    draw_heart(s2, 1200, 980, s=13, angle=14)
    draw_star(s2, 1120, 540, s=13, angle=25)
    draw_spark(s2, 640, 980, s=11)
    draw_star(s2, 1820, 980, s=12, angle=-12)
    draw_spark(s2, 980, 980, s=9)


def draw_decor_s3(s3, box_ch, box_dr):
    f_title = load_font(FONT_TITLE, 120)
    f_lab3 = load_font(FONT_LABEL, 44)
    draw_hand_text(s3, "现在的你", (700, 36), f_title, angle=-2.6, jitter=2.3, seed=17)
    draw_hand_text(s3, "认真生活～闪闪发光", (60, 910), load_font(FONT_TITLE, 64), angle=-3.2, jitter=1.6, seed=18)
    draw_hand_text(s3, "闪光日常", (760, 250), f_lab3, angle=10, jitter=1.3, seed=20)
    draw_arrow(s3, (860, 320), (box_ch[2] - 30, box_ch[1] + 70), dotted=True, seed=8)
    draw_heart(s3, 1080, 150, s=18, angle=-12)
    draw_heart(s3, 520, 140, s=13, angle=22)
    draw_star(s3, 1740, 70, s=15, angle=18)
    draw_spark(s3, 980, 620, s=13)
    draw_star(s3, 1780, 980, s=13, angle=-8)
    draw_heart(s3, 1100, 980, s=12, angle=16)


def build_layered_pptx(gif_paths: list[Path], out_path: Path):
    """GIF slides + layered static slides with OOXML teeter on stickers."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def px_to_left(x):
        return Inches(x / 1920 * 13.333333)

    def px_to_top(y):
        return Inches(y / 1080 * 7.5)

    def px_to_w(w):
        return Inches(w / 1920 * 13.333333)

    def px_to_h(h):
        return Inches(h / 1080 * 7.5)

    for gif in gif_paths:
        slide = prs.slides.add_slide(blank)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
        shape.line.fill.background()
        slide.shapes.add_picture(str(gif), Inches(0), Inches(0), prs.slide_width, prs.slide_height)

    stickers = prepare_stickers()
    layer_dir = ROOT / "layers"
    layer_dir.mkdir(exist_ok=True)

    # Slide 1 layered
    s1_child = rotate_keep(stickers["child"], -5.0)
    p_child = layer_dir / "s1_child.png"
    s1_child.save(p_child)
    cw, ch = s1_child.size
    s1_pos = (620 - cw // 2, 545 - ch // 2)
    box_c = (s1_pos[0], s1_pos[1], s1_pos[0] + cw, s1_pos[1] + ch)
    decor1 = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    draw_decor_s1(decor1, box_c)
    p_decor1 = layer_dir / "s1_decor.png"
    decor1.save(p_decor1)

    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0, 0, 0)
    bg.line.fill.background()
    pic1 = slide.shapes.add_picture(str(p_child), px_to_left(s1_pos[0]), px_to_top(s1_pos[1]), px_to_w(cw), px_to_h(ch))
    slide.shapes.add_picture(str(p_decor1), Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    inject_teeter_timing(slide, [pic1.shape_id])

    # Slide 2 layered
    s2_cake = rotate_keep(stickers["cake"], -7.5)
    s2_art = rotate_keep(stickers["artist"], 6.8)
    s2_dress = rotate_keep(stickers["dress"], 3.5)
    p_cake, p_art, p_dress = layer_dir / "s2_cake.png", layer_dir / "s2_art.png", layer_dir / "s2_dress.png"
    s2_cake.save(p_cake)
    s2_art.save(p_art)
    s2_dress.save(p_dress)
    pos_cake, pos_art, pos_dress = (40, 300), (680, 95), (1310, 170)
    box_cake = (40, 300, 40 + s2_cake.size[0], 300 + s2_cake.size[1])
    box_art = (680, 95, 680 + s2_art.size[0], 95 + s2_art.size[1])
    box_dress = (1310, 170, 1310 + s2_dress.size[0], 170 + s2_dress.size[1])
    decor2 = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    draw_decor_s2(decor2, box_cake, box_art, box_dress)
    p_decor2 = layer_dir / "s2_decor.png"
    decor2.save(p_decor2)

    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0, 0, 0)
    bg.line.fill.background()
    pics = []
    for path, pos, im in [(p_cake, pos_cake, s2_cake), (p_art, pos_art, s2_art), (p_dress, pos_dress, s2_dress)]:
        pic = slide.shapes.add_picture(
            str(path), px_to_left(pos[0]), px_to_top(pos[1]), px_to_w(im.size[0]), px_to_h(im.size[1])
        )
        pics.append(pic.shape_id)
    slide.shapes.add_picture(str(p_decor2), Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    inject_teeter_timing(slide, pics)

    # Slide 3 layered
    s3_chili = rotate_keep(stickers["chili"], 6.2)
    s3_dress = rotate_keep(stickers["dress3"], -5.8)
    p_chili, p_dress3 = layer_dir / "s3_chili.png", layer_dir / "s3_dress.png"
    s3_chili.save(p_chili)
    s3_dress.save(p_dress3)
    pos_chili, pos_dress3 = (70, 200), (1200, 80)
    box_ch = (70, 200, 70 + s3_chili.size[0], 200 + s3_chili.size[1])
    box_dr = (1200, 80, 1200 + s3_dress.size[0], 80 + s3_dress.size[1])
    decor3 = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    draw_decor_s3(decor3, box_ch, box_dr)
    p_decor3 = layer_dir / "s3_decor.png"
    decor3.save(p_decor3)

    slide = prs.slides.add_slide(blank)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0, 0, 0)
    bg.line.fill.background()
    pics = []
    for path, pos, im in [(p_chili, pos_chili, s3_chili), (p_dress3, pos_dress3, s3_dress)]:
        pic = slide.shapes.add_picture(
            str(path), px_to_left(pos[0]), px_to_top(pos[1]), px_to_w(im.size[0]), px_to_h(im.size[1])
        )
        pics.append(pic.shape_id)
    slide.shapes.add_picture(str(p_decor3), Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    inject_teeter_timing(slide, pics)

    prs.save(str(out_path))
    print(f"wrote {out_path} ({out_path.stat().st_size / 1e6:.2f} MB)")


def main():
    print("Preparing stickers...")
    stickers = prepare_stickers()

    builders = [
        ("slide1_jitter.gif", build_slide1_frame),
        ("slide2_jitter.gif", build_slide2_frame),
        ("slide3_jitter.gif", build_slide3_frame),
    ]
    gif_paths = []
    for name, builder in builders:
        print(f"Rendering {name} ({N_FRAMES} frames @ {FPS}fps)...")
        frames = []
        for i in range(N_FRAMES):
            t = i / N_FRAMES
            frames.append(builder(stickers, t))
        path = ROOT / name
        save_gif(frames, path)
        gif_paths.append(path)
        mp4 = ROOT / name.replace(".gif", ".mp4")
        try:
            small_frames = [f.resize((1280, 720), Image.Resampling.LANCZOS) for f in frames]
            save_mp4(small_frames, mp4)
        except Exception as e:
            print("mp4 skip:", e)

    print("Building animated PPTX...")
    pptx_path = ROOT / "preview_animated.pptx"
    build_layered_pptx(gif_paths, pptx_path)
    print("DONE (no zip — share files individually)")


if __name__ == "__main__":
    main()
