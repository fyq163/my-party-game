#!/usr/bin/env python3
"""Build a 16:9 editable PPTX: movable sticker pictures + native text boxes.

Unlike generate_slides.py (flattened PNG slides), this keeps each sticker as its
own picture and titles/labels as PowerPoint text so hosts can tweak copy and
placement in Keynote / PowerPoint / LibreOffice after generation.

Requires cutouts in assets/ (same keys as generate_slides.py). Decor PNGs in
assets/ are used for hearts, stars, and dotted arrows.
"""

from __future__ import annotations

import argparse
import math
import os
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from generate_slides import (
    ASSETS,
    H,
    ROOT,
    W,
    load_cutouts,
    make_sticker,
    rotate_keep,
)

SLIDE_W_IN = 13.333333
SLIDE_H_IN = 7.5

FONT_TITLE_NAME = "Zhi Mang Xing"
FONT_LABEL_NAME = "Ma Shan Zheng"
PINK_RGB = RGBColor(255, 45, 138)
WHITE_RGB = RGBColor(255, 255, 255)


def px_left(x: float):
    return Inches(x / W * SLIDE_W_IN)


def px_top(y: float):
    return Inches(y / H * SLIDE_H_IN)


def px_w(w: float):
    return Inches(w / W * SLIDE_W_IN)


def px_h(h: float):
    return Inches(h / H * SLIDE_H_IN)


def px_to_pt(px: float) -> float:
    """1920×1080 on a 13.333″×7.5″ slide is 144 dpi; pt = px * 72 / 144."""
    return px * 0.5


def add_black_bg(slide, prs):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0, 0, 0)
    shape.line.fill.background()
    # Keep the fill behind pictures/text.
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return shape


def set_east_asian_font(run, name: str):
    rPr = run._r.get_or_add_rPr()
    for tag, attr_name in (("a:latin", name), ("a:ea", name), ("a:cs", name)):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", attr_name)


def add_text(
    slide,
    text: str,
    xy: tuple[float, float],
    size_px: float,
    font_name: str,
    color=WHITE_RGB,
    angle: float = 0.0,
    box: tuple[float, float] = (520, 160),
    align=PP_ALIGN.LEFT,
):
    x, y = xy
    bw, bh = box
    tx = slide.shapes.add_textbox(px_left(x), px_top(y), px_w(bw), px_h(bh))
    tf = tx.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(px_to_pt(size_px))
    run.font.color.rgb = color
    set_east_asian_font(run, font_name)
    if abs(angle) > 0.05:
        tx.rotation = angle
    return tx


def add_picture(slide, path: Path, xy: tuple[float, float], size: tuple[int, int]):
    x, y = xy
    w, h = size
    return slide.shapes.add_picture(str(path), px_left(x), px_top(y), px_w(w), px_h(h))


def place_sticker(slide, layer_dir: Path, name: str, im, xy: tuple[float, float], anchor="lt"):
    path = layer_dir / f"{name}.png"
    im.save(path)
    w, h = im.size
    x, y = xy
    if anchor == "center":
        x -= w / 2
        y -= h / 2
    pic = add_picture(slide, path, (x, y), (w, h))
    box = (x, y, x + w, y + h)
    return pic, box


def place_decor(slide, filename: str, cx: float, cy: float, scale: float = 2.4, angle: float = 0.0):
    src = ASSETS / filename
    if not src.exists():
        return None
    from PIL import Image

    im = Image.open(src).convert("RGBA")
    w, h = im.size
    w, h = int(w * scale), int(h * scale)
    pic = add_picture(slide, src, (cx - w / 2, cy - h / 2), (w, h))
    if abs(angle) > 0.05:
        pic.rotation = angle
    return pic


def place_arrow(slide, start: tuple[float, float], end: tuple[float, float]):
    """Place the dotted-arrow PNG, rotated toward `end`."""
    src = ASSETS / "arrow_dotted.png"
    if not src.exists():
        return None
    x0, y0 = start
    x1, y1 = end
    mx, my = (x0 + x1) / 2, (y0 + y1) / 2
    dx, dy = x1 - x0, y1 - y0
    dist = math.hypot(dx, dy)
    # PNG arrow points roughly left/down; rotate to the segment.
    native_deg = math.degrees(math.atan2(40, -180))  # approx of the asset
    target_deg = math.degrees(math.atan2(dy, dx))
    angle = target_deg - native_deg
    w = max(120, min(280, dist * 0.55))
    h = w * 120 / 220
    pic = add_picture(slide, src, (mx - w / 2, my - h / 2), (w, h))
    pic.rotation = angle
    return pic


def build_editable_pptx(
    out_path: Path,
    child_src: Path | None = None,
    collage_src: Path | None = None,
):
    cuts = load_cutouts(child_src, collage_src)
    layer_dir = ROOT / "layers"
    layer_dir.mkdir(exist_ok=True)

    stickers = {
        "child": rotate_keep(make_sticker(cuts["child"], border=20, jitter_amp=6.2, seed=21, max_h=880), -5.0),
        "cake": rotate_keep(make_sticker(cuts["cake"], border=16, jitter_amp=5.4, seed=3, max_h=500, max_w=640), -7.5),
        "artist": rotate_keep(make_sticker(cuts["artist"], border=15, jitter_amp=5.1, seed=8, max_h=390, max_w=560), 6.8),
        "dress": rotate_keep(make_sticker(cuts["dress"], border=17, jitter_amp=5.8, seed=12, max_h=740), 3.5),
        "dress3": rotate_keep(make_sticker(cuts["dress"], border=18, jitter_amp=5.7, seed=19, max_h=900), -5.8),
        "chili": rotate_keep(make_sticker(cuts["chili"], border=16, jitter_amp=5.4, seed=22, max_h=540, max_w=760), 6.2),
    }

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    blank = prs.slide_layouts[6]

    # ----- Slide 1: 小时候 -----
    s1 = prs.slides.add_slide(blank)
    add_black_bg(s1, prs)
    _, box_c = place_sticker(s1, layer_dir, "edit_s1_child", stickers["child"], (620, 545), anchor="center")
    add_text(s1, "小时候～", (1180, 40), 120, FONT_TITLE_NAME, angle=-3.5, box=(620, 180))
    add_text(s1, "～", (1580, 210), 58, FONT_TITLE_NAME, color=PINK_RGB, angle=14, box=(160, 120))
    add_text(s1, "小小的你", (1320, 700), 50, FONT_LABEL_NAME, angle=7.0, box=(360, 120))
    place_arrow(s1, (1310, 750), (box_c[2] - 40, (box_c[1] + box_c[3]) / 2 + 30))
    place_decor(s1, "decor_heart.png", 1160, 300, scale=2.6, angle=-16)
    place_decor(s1, "decor_star.png", 1700, 880, scale=2.4, angle=22)
    place_decor(s1, "decor_heart.png", 1640, 500, scale=2.0, angle=18)

    # ----- Slide 2: 长大啦 -----
    s2 = prs.slides.add_slide(blank)
    add_black_bg(s2, prs)
    _, box_cake = place_sticker(s2, layer_dir, "edit_s2_cake", stickers["cake"], (40, 300))
    _, box_art = place_sticker(s2, layer_dir, "edit_s2_art", stickers["artist"], (680, 95))
    _, box_dress = place_sticker(s2, layer_dir, "edit_s2_dress", stickers["dress"], (1310, 170))
    add_text(s2, "长大啦！！", (36, 28), 112, FONT_TITLE_NAME, angle=-5.8, box=(720, 180))
    add_text(s2, "吃蛋糕", (80, 920), 46, FONT_LABEL_NAME, angle=-6, box=(280, 110))
    add_text(s2, "画画", (1040, 70), 46, FONT_LABEL_NAME, angle=8, box=(220, 110))
    add_text(s2, "穿搭", (1688, 90), 46, FONT_LABEL_NAME, angle=8, box=(220, 110))
    place_arrow(s2, (220, 910), ((box_cake[0] + box_cake[2]) / 2, box_cake[3] - 24))
    place_arrow(s2, (1100, 145), (box_art[2] - 40, box_art[1] + 70))
    place_arrow(s2, (1740, 165), (box_dress[0] + 90, box_dress[1] + 50))
    place_decor(s2, "decor_heart.png", 620, 180, scale=2.4, angle=-22)
    place_decor(s2, "decor_heart.png", 1200, 980, scale=2.0, angle=14)
    place_decor(s2, "decor_star.png", 1120, 540, scale=2.2, angle=25)
    place_decor(s2, "decor_star.png", 1820, 980, scale=2.0, angle=-12)

    # ----- Slide 3: 现在的你 -----
    s3 = prs.slides.add_slide(blank)
    add_black_bg(s3, prs)
    _, box_ch = place_sticker(s3, layer_dir, "edit_s3_chili", stickers["chili"], (70, 200))
    place_sticker(s3, layer_dir, "edit_s3_dress", stickers["dress3"], (1200, 80))
    add_text(s3, "现在的你", (700, 36), 120, FONT_TITLE_NAME, angle=-2.6, box=(720, 180))
    add_text(s3, "认真生活～闪闪发光", (60, 910), 64, FONT_TITLE_NAME, angle=-3.2, box=(1100, 140))
    add_text(s3, "闪光日常", (760, 250), 44, FONT_LABEL_NAME, angle=10, box=(320, 110))
    place_arrow(s3, (860, 320), (box_ch[2] - 30, box_ch[1] + 70))
    place_decor(s3, "decor_heart.png", 1080, 150, scale=2.8, angle=-12)
    place_decor(s3, "decor_heart.png", 520, 140, scale=2.0, angle=22)
    place_decor(s3, "decor_star.png", 1740, 70, scale=2.4, angle=18)
    place_decor(s3, "decor_star.png", 1780, 980, scale=2.2, angle=-8)
    place_decor(s3, "decor_heart.png", 1100, 980, scale=2.0, angle=16)

    prs.save(str(out_path))
    print(f"wrote {out_path} ({out_path.stat().st_size / 1e6:.2f} MB)")
    print("Note: native text uses Zhi Mang Xing / Ma Shan Zheng; install those")
    print("fonts on the editing machine, or use generate_slides.py for baked-in PNG slides.")


def main():
    parser = argparse.ArgumentParser(description="Build an editable black-collage 16:9 PPTX.")
    parser.add_argument("--child", default=os.environ.get("CHILD_IMG"), help="optional source image for child cutout")
    parser.add_argument("--collage", default=os.environ.get("COLLAGE_IMG"), help="optional source collage to re-extract")
    parser.add_argument(
        "--out",
        default=str(ROOT / "editable.pptx"),
        help="output .pptx path (copy into the rundown ppt path after editing)",
    )
    args = parser.parse_args()
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    build_editable_pptx(
        out,
        Path(args.child) if args.child else None,
        Path(args.collage) if args.collage else None,
    )


if __name__ == "__main__":
    main()
