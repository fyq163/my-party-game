#!/usr/bin/env python3
"""Nostalgic black scrapbook / vlog collage slides.

Paths are relative to this toolkit folder. Prefers existing cutouts in assets/;
optional source images can be provided via env vars or CLI args to re-extract.
"""

from __future__ import annotations

import math
import os
import sys
from pathlib import Path

import cv2
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
FONTS = ROOT / "fonts"
ASSETS.mkdir(parents=True, exist_ok=True)

PINK = (255, 45, 138, 255)
W, H = 1920, 1080

FONT_TITLE = str(FONTS / "ZhiMangXing-Regular.ttf")
FONT_LABEL = str(FONTS / "MaShanZheng-Regular.ttf")
FONT_ALT = str(FONTS / "LiuJianMaoCao-Regular.ttf")


def load_font(path: str, size: int) -> ImageFont.FreeTypeFont:
    try:
        return ImageFont.truetype(path, size=size)
    except Exception:
        return ImageFont.truetype(FONT_ALT, size=size)


def trim_rgba(im: Image.Image, pad: int = 2) -> Image.Image:
    arr = np.array(im)
    if arr.shape[2] < 4:
        return im
    ys, xs = np.where(arr[:, :, 3] > 8)
    if len(xs) == 0:
        return im
    x0, x1 = max(0, int(xs.min()) - pad), min(arr.shape[1], int(xs.max()) + pad + 1)
    y0, y1 = max(0, int(ys.min()) - pad), min(arr.shape[0], int(ys.max()) + pad + 1)
    return im.crop((x0, y0, x1, y1))


def fill_interior_holes(mask: np.ndarray, max_area: int | None = None) -> np.ndarray:
    h, w = mask.shape
    inv = cv2.bitwise_not(mask)
    n, labels, stats, _ = cv2.connectedComponentsWithStats(inv, 8)
    out = mask.copy()
    for i in range(1, n):
        ys, xs = np.where(labels == i)
        if ys.size == 0:
            continue
        if ys.min() == 0 or xs.min() == 0 or ys.max() == h - 1 or xs.max() == w - 1:
            continue
        if max_area is not None and stats[i, cv2.CC_STAT_AREA] > max_area:
            continue
        out[labels == i] = 255
    return out


def defringe_white(rgb: np.ndarray, alpha: np.ndarray, rounds: int = 4) -> np.ndarray:
    a = alpha.copy()
    dist = np.linalg.norm(rgb.astype(np.int16) - 255, axis=2)
    near_white = dist < 28
    k = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (3, 3))
    for _ in range(rounds):
        empty = (a < 12).astype(np.uint8)
        ring = cv2.dilate(empty, k) > 0
        a[ring & near_white] = 0
    return a


def black_bg_to_alpha(path: Path) -> Image.Image:
    rgb = cv2.cvtColor(cv2.imread(str(path)), cv2.COLOR_BGR2RGB)
    mx = rgb.max(axis=2)
    ink = (mx > 10).astype(np.uint8) * 255
    ink = cv2.morphologyEx(
        ink, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (7, 7)), iterations=2
    )
    n, labels, stats, _ = cv2.connectedComponentsWithStats(ink, 8)
    idx = 1 + np.argmax(stats[1:, cv2.CC_STAT_AREA])
    m = np.where(labels == idx, 255, 0).astype(np.uint8)
    m = fill_interior_holes(m, max_area=4000)
    m = cv2.GaussianBlur(m, (5, 5), 0)
    rgba = np.dstack([rgb, m])
    return trim_rgba(Image.fromarray(rgba, "RGBA"), 1)


def extract_collage_cutouts(path: Path) -> dict[str, Image.Image]:
    rgb = cv2.cvtColor(cv2.imread(str(path)), cv2.COLOR_BGR2RGB)
    h, w = rgb.shape[:2]
    dist = np.linalg.norm(rgb.astype(np.int16) - 255, axis=2)
    ink = (dist > 16).astype(np.uint8) * 255
    closed = cv2.morphologyEx(
        ink, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (7, 7)), iterations=2
    )
    n, labels, stats, cents = cv2.connectedComponentsWithStats(closed, 8)
    items = []
    for i in range(1, n):
        if stats[i, cv2.CC_STAT_AREA] < 2000:
            continue
        items.append((int(stats[i, 4]), i, stats[i], cents[i]))
    items.sort(reverse=True)
    items = items[:4]
    other_map = (labels > 0).astype(np.uint8)

    named: dict[str, Image.Image] = {}
    for area, i, st, ct in items:
        cx, cy = float(ct[0]), float(ct[1])
        if cy < h * 0.38 and 0.28 * w < cx < 0.72 * w:
            key = "artist"
        elif cx > 0.68 * w:
            key = "dress"
        elif cy > 0.52 * h and cx > 0.38 * w:
            key = "chili"
        else:
            key = "cake"

        m = np.where(labels == i, 255, 0).astype(np.uint8)
        if key == "dress":
            m = cv2.morphologyEx(
                m, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (21, 21)), iterations=2
            )
            m = fill_interior_holes(m, max_area=None)
        else:
            m = cv2.morphologyEx(
                m, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (9, 9)), iterations=1
            )
            m = fill_interior_holes(m, max_area=8000)
        m[(other_map > 0) & (labels != i)] = 0
        m = cv2.dilate(m, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (3, 3)))
        m[(other_map > 0) & (labels != i)] = 0

        x, y, cw, ch = map(int, st[:4])
        pad = 10
        x0, y0 = max(0, x - pad), max(0, y - pad)
        x1, y1 = min(w, x + cw + pad), min(h, y + ch + pad)
        crop = rgb[y0:y1, x0:x1]
        alpha = m[y0:y1, x0:x1]
        alpha = defringe_white(crop, alpha, rounds=5)
        alpha = cv2.GaussianBlur(alpha, (3, 3), 0)
        im = trim_rgba(Image.fromarray(np.dstack([crop, alpha]), "RGBA"), 1)
        named[key] = im
        im.save(ASSETS / f"cutout_{key}.png")
        print(f"cutout {key}: {im.size} area={area}")
    return named


def resample_contour(pts: np.ndarray, n: int = 260) -> np.ndarray:
    pts = pts.reshape(-1, 2).astype(np.float64)
    if len(pts) < 3:
        return pts
    if np.linalg.norm(pts[0] - pts[-1]) > 1:
        pts = np.vstack([pts, pts[0]])
    d = np.sqrt(((pts[1:] - pts[:-1]) ** 2).sum(axis=1))
    u = np.concatenate([[0], np.cumsum(d)])
    if u[-1] < 1e-6:
        return pts[:-1]
    u /= u[-1]
    t = np.linspace(0, 1, n, endpoint=False)
    return np.stack([np.interp(t, u, pts[:, 0]), np.interp(t, u, pts[:, 1])], axis=1)


def jitter_contour(pts: np.ndarray, amp: float, seed: int) -> np.ndarray:
    rng = np.random.RandomState(seed)
    n = len(pts)
    t = np.linspace(0, 2 * np.pi, n, endpoint=False)
    phase = rng.uniform(0, 2 * np.pi)
    wave = (
        0.55 * np.sin(2 * t + phase)
        + 0.30 * np.sin(5 * t + phase * 1.7)
        + 0.15 * np.sin(9 * t + phase * 0.4)
        + 0.10 * rng.randn(n)
    )
    k = 11
    wave = np.convolve(np.pad(wave, k, mode="wrap"), np.ones(k) / k, mode="valid")[:n]
    nxt = np.roll(pts, -1, axis=0)
    prv = np.roll(pts, 1, axis=0)
    tang = nxt - prv
    norm = np.stack([-tang[:, 1], tang[:, 0]], axis=1)
    ln = np.linalg.norm(norm, axis=1, keepdims=True) + 1e-6
    norm = norm / ln
    return pts + norm * (wave * amp)[:, None]


def make_sticker(
    src: Image.Image,
    border: int = 16,
    jitter_amp: float = 5.5,
    seed: int = 7,
    max_h: int | None = None,
    max_w: int | None = None,
) -> Image.Image:
    im = src.convert("RGBA")
    w0, h0 = im.size
    s = 1.0
    if max_h:
        s = min(s, max_h / h0)
    if max_w:
        s = min(s, max_w / w0)
    if s != 1.0:
        im = im.resize((max(1, int(w0 * s)), max(1, int(h0 * s))), Image.Resampling.LANCZOS)

    arr = np.array(im)
    alpha = arr[:, :, 3]
    m = (alpha > 20).astype(np.uint8) * 255
    m = cv2.morphologyEx(m, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (5, 5)))

    pad = border + 24
    mpad = cv2.copyMakeBorder(m, pad, pad, pad, pad, cv2.BORDER_CONSTANT, value=0)
    rgbpad = cv2.copyMakeBorder(arr[:, :, :3], pad, pad, pad, pad, cv2.BORDER_CONSTANT, value=0)
    apad = cv2.copyMakeBorder(alpha, pad, pad, pad, pad, cv2.BORDER_CONSTANT, value=0)

    ksz = max(5, int(border * 1.9) | 1)
    dilated = cv2.dilate(mpad, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (ksz, ksz)))
    contours, _ = cv2.findContours(dilated, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_NONE)
    if not contours:
        return Image.fromarray(np.dstack([rgbpad, apad]), "RGBA")
    cnt = max(contours, key=cv2.contourArea)
    pts = jitter_contour(resample_contour(cnt, n=300), amp=jitter_amp, seed=seed)
    pts_i = np.round(pts).astype(np.int32)

    pink_layer = np.zeros((*dilated.shape, 4), np.uint8)
    cv2.fillPoly(pink_layer, [pts_i], PINK)
    cv2.polylines(pink_layer, [pts_i], True, PINK, thickness=max(3, border // 3), lineType=cv2.LINE_AA)
    pts2 = jitter_contour(resample_contour(cnt, n=300), amp=jitter_amp * 0.4, seed=seed + 11)
    cv2.polylines(
        pink_layer,
        [np.round(pts2).astype(np.int32)],
        True,
        (255, 90, 165, 200),
        thickness=max(2, border // 5),
        lineType=cv2.LINE_AA,
    )

    out = pink_layer.copy()
    a = apad.astype(np.float32) / 255.0
    for c in range(3):
        out[:, :, c] = np.clip(rgbpad[:, :, c] * a + out[:, :, c] * (1 - a), 0, 255).astype(np.uint8)
    out[:, :, 3] = np.clip(apad.astype(np.int16) + out[:, :, 3].astype(np.int16), 0, 255).astype(np.uint8)
    strong = apad > 40
    out[strong, :3] = rgbpad[strong]
    out[strong, 3] = np.maximum(out[strong, 3], apad[strong])
    return trim_rgba(Image.fromarray(out, "RGBA"), 1)


def rotate_keep(im: Image.Image, deg: float) -> Image.Image:
    return im.rotate(deg, expand=True, resample=Image.Resampling.BICUBIC)


def paste(canvas: Image.Image, sticker: Image.Image, xy: tuple[int, int], anchor: str = "lt"):
    x, y = xy
    w, h = sticker.size
    if anchor == "center":
        x -= w // 2
        y -= h // 2
    else:
        if "c" in anchor and "l" not in anchor and "r" not in anchor:
            x -= w // 2
        if anchor.endswith("r") or anchor.startswith("r"):
            x -= w
        if "b" in anchor:
            y -= h
    x = int(max(4, min(x, W - w - 4)))
    y = int(max(4, min(y, H - h - 4)))
    canvas.alpha_composite(sticker, (x, y))
    return (x, y, x + w, y + h)


def draw_hand_text(
    canvas: Image.Image,
    text: str,
    xy: tuple[int, int],
    font: ImageFont.FreeTypeFont,
    fill=(255, 255, 255, 255),
    angle: float = 0.0,
    jitter: float = 1.8,
    seed: int = 1,
    char_space: float = 1.0,
):
    rng = np.random.RandomState(seed)
    tmp = Image.new("RGBA", (8, 8), (0, 0, 0, 0))
    td = ImageDraw.Draw(tmp)
    widths, heights = [], []
    for ch in text:
        bbox = td.textbbox((0, 0), ch, font=font)
        widths.append(max(1, bbox[2] - bbox[0]))
        heights.append(max(1, bbox[3] - bbox[1]))
    total_w = int(sum(widths) * char_space + 24)
    total_h = int(max(heights) + 36)
    layer = Image.new("RGBA", (total_w + 50, total_h + 50), (0, 0, 0, 0))
    x = 18
    base_y = 18
    for i, ch in enumerate(text):
        ox = rng.uniform(-jitter, jitter)
        oy = rng.uniform(-jitter * 1.15, jitter * 1.15)
        rot = rng.uniform(-4.2, 4.2)
        cb = td.textbbox((0, 0), ch, font=font)
        cw, chh = cb[2] - cb[0] + 16, cb[3] - cb[1] + 20
        ch_im = Image.new("RGBA", (max(cw, 8), max(chh, 8) + 12), (0, 0, 0, 0))
        cd = ImageDraw.Draw(ch_im)
        cd.text((5, 3), ch, font=font, fill=fill)
        cd.text((4, 2), ch, font=font, fill=fill)
        ch_im = ch_im.rotate(rot, expand=True, resample=Image.Resampling.BICUBIC)
        layer.alpha_composite(ch_im, (int(x + ox), int(base_y + oy)))
        x += widths[i] * char_space
    if abs(angle) > 0.05:
        layer = layer.rotate(angle, expand=True, resample=Image.Resampling.BICUBIC)
    layer = trim_rgba(layer, 1)
    x0, y0 = int(xy[0]), int(xy[1])
    x0 = max(0, min(x0, W - 4))
    y0 = max(0, min(y0, H - 4))
    canvas.alpha_composite(layer, (x0, y0))
    return (x0, y0, x0 + layer.size[0], y0 + layer.size[1])


def draw_arrow(canvas, start, end, color=(255, 255, 255, 255), dotted=True, width=3, seed=3):
    rng = np.random.RandomState(seed)
    x0, y0 = start
    x1, y1 = end
    mx, my = (x0 + x1) / 2, (y0 + y1) / 2
    nx, ny = -(y1 - y0), (x1 - x0)
    ln = math.hypot(nx, ny) + 1e-6
    nx, ny = nx / ln, ny / ln
    bulge = rng.uniform(0.10, 0.22) * math.hypot(x1 - x0, y1 - y0) * rng.choice([-1, 1])
    cx, cy = mx + nx * bulge, my + ny * bulge
    n = 40
    pts = []
    for i in range(n + 1):
        t = i / n
        x = (1 - t) ** 2 * x0 + 2 * (1 - t) * t * cx + t**2 * x1
        y = (1 - t) ** 2 * y0 + 2 * (1 - t) * t * cy + t**2 * y1
        pts.append((x + rng.uniform(-0.5, 0.5), y + rng.uniform(-0.5, 0.5)))
    d = ImageDraw.Draw(canvas)
    if dotted:
        for i in range(0, len(pts) - 1, 2):
            d.line([pts[i], pts[i + 1]], fill=color, width=width)
    else:
        d.line(pts, fill=color, width=width)
    ang = math.atan2(pts[-1][1] - pts[-3][1], pts[-1][0] - pts[-3][0])
    ah = 18
    p1 = (end[0] - ah * math.cos(ang - 0.48), end[1] - ah * math.sin(ang - 0.48))
    p2 = (end[0] - ah * math.cos(ang + 0.48), end[1] - ah * math.sin(ang + 0.48))
    d.line([p1, end], fill=color, width=width + 1)
    d.line([p2, end], fill=color, width=width + 1)


def draw_heart(canvas, cx, cy, s=18, color=PINK, angle=0):
    layer = Image.new("RGBA", (s * 4, s * 4), (0, 0, 0, 0))
    d = ImageDraw.Draw(layer)
    pts = []
    for t in np.linspace(0, 2 * math.pi, 90):
        x = 16 * math.sin(t) ** 3
        y = -(13 * math.cos(t) - 5 * math.cos(2 * t) - 2 * math.cos(3 * t) - math.cos(4 * t))
        pts.append((s * 2 + x * (s / 16.0), s * 2 + y * (s / 16.0)))
    d.line(pts + [pts[0]], fill=color, width=max(3, s // 7))
    if abs(angle) > 0.1:
        layer = layer.rotate(angle, expand=True, resample=Image.Resampling.BICUBIC)
    canvas.alpha_composite(layer, (cx - layer.size[0] // 2, cy - layer.size[1] // 2))


def draw_star(canvas, cx, cy, s=14, color=PINK, angle=15):
    layer = Image.new("RGBA", (s * 4, s * 4), (0, 0, 0, 0))
    d = ImageDraw.Draw(layer)
    pts = []
    for i in range(10):
        ang = math.radians(-90 + i * 36)
        r = s if i % 2 == 0 else s * 0.42
        pts.append((s * 2 + r * math.cos(ang), s * 2 + r * math.sin(ang)))
    d.line(pts + [pts[0]], fill=color, width=max(2, s // 6))
    layer = layer.rotate(angle, expand=True, resample=Image.Resampling.BICUBIC)
    canvas.alpha_composite(layer, (cx - layer.size[0] // 2, cy - layer.size[1] // 2))


def draw_spark(canvas, cx, cy, s=12, color=PINK):
    d = ImageDraw.Draw(canvas)
    d.line([(cx - s, cy), (cx + s, cy)], fill=color, width=3)
    d.line([(cx, cy - s), (cx, cy + s)], fill=color, width=3)
    d.line([(cx - s * 0.55, cy - s * 0.55), (cx + s * 0.55, cy + s * 0.55)], fill=color, width=2)
    d.line([(cx - s * 0.55, cy + s * 0.55), (cx + s * 0.55, cy - s * 0.55)], fill=color, width=2)


def new_canvas() -> Image.Image:
    return Image.new("RGBA", (W, H), (0, 0, 0, 255))


def flatten(im: Image.Image) -> Image.Image:
    bg = Image.new("RGB", im.size, (0, 0, 0))
    bg.paste(im, mask=im.split()[-1])
    return bg


def load_cutouts(child_src: Path | None = None, collage_src: Path | None = None) -> dict[str, Image.Image]:
    """Load cutouts from assets/; optionally re-extract from source images."""
    need = ["child", "cake", "artist", "dress", "chili"]
    cuts: dict[str, Image.Image] = {}

    if child_src and child_src.exists():
        print(f"extracting child from {child_src}...")
        child = black_bg_to_alpha(child_src)
        child.save(ASSETS / "cutout_child.png")
        cuts["child"] = child
    if collage_src and collage_src.exists():
        print(f"extracting collage from {collage_src}...")
        cuts.update(extract_collage_cutouts(collage_src))

    for key in need:
        if key in cuts:
            continue
        p = ASSETS / f"cutout_{key}.png"
        if not p.exists():
            raise FileNotFoundError(
                f"Missing {p}. Place cutout PNGs in assets/ or pass source images "
                f"(CHILD_IMG / COLLAGE_IMG env, or --child / --collage)."
            )
        cuts[key] = Image.open(p).convert("RGBA")
        print(f"loaded {p.name}: {cuts[key].size}")
    return cuts


def build_slides(child_src: Path | None = None, collage_src: Path | None = None):
    cuts = load_cutouts(child_src, collage_src)

    f_title = load_font(FONT_TITLE, 120)
    f_title2 = load_font(FONT_TITLE, 112)
    f_label = load_font(FONT_LABEL, 50)

    # ---------- SLIDE 1 ----------
    print("slide 1...")
    s1 = new_canvas()
    st_child = make_sticker(cuts["child"], border=20, jitter_amp=6.2, seed=21, max_h=880)
    st_child = rotate_keep(st_child, -5.0)
    box_c = paste(s1, st_child, (620, 545), anchor="center")

    draw_hand_text(s1, "小时候～", (1180, 40), f_title, angle=-3.5, jitter=2.5, seed=4)
    draw_hand_text(s1, "～", (1580, 210), load_font(FONT_TITLE, 58), fill=PINK, angle=14, jitter=1.2, seed=9)
    draw_hand_text(s1, "小小的你", (1320, 700), f_label, angle=7.0, jitter=1.6, seed=6)
    draw_arrow(
        s1,
        (1310, 750),
        (box_c[2] - 40, (box_c[1] + box_c[3]) // 2 + 30),
        dotted=True,
        width=3,
        seed=2,
    )
    draw_heart(s1, 1160, 300, s=17, angle=-16)
    draw_star(s1, 1700, 880, s=14, angle=22)
    draw_spark(s1, 220, 160, s=12)
    draw_heart(s1, 1640, 500, s=13, angle=18)
    draw_spark(s1, 1080, 980, s=9)

    p1 = ROOT / "slide1_childhood.png"
    flatten(s1).save(p1, "PNG")
    print("wrote", p1)

    # ---------- SLIDE 2 ----------
    print("slide 2...")
    s2 = new_canvas()
    st_cake = make_sticker(cuts["cake"], border=16, jitter_amp=5.4, seed=3, max_h=500, max_w=640)
    st_cake = rotate_keep(st_cake, -7.5)
    st_art = make_sticker(cuts["artist"], border=15, jitter_amp=5.1, seed=8, max_h=390, max_w=560)
    st_art = rotate_keep(st_art, 6.8)
    st_dress = make_sticker(cuts["dress"], border=17, jitter_amp=5.8, seed=12, max_h=740)
    st_dress = rotate_keep(st_dress, 3.5)

    box_cake = paste(s2, st_cake, (40, 300), anchor="lt")
    box_art = paste(s2, st_art, (680, 95), anchor="lt")
    box_dress = paste(s2, st_dress, (1310, 170), anchor="lt")

    draw_hand_text(s2, "长大啦！！", (36, 28), f_title2, angle=-5.8, jitter=2.3, seed=11)
    f_lab = load_font(FONT_LABEL, 46)
    draw_hand_text(s2, "吃蛋糕", (80, 920), f_lab, angle=-6, jitter=1.4, seed=13)
    draw_arrow(s2, (220, 910), ((box_cake[0] + box_cake[2]) // 2, box_cake[3] - 24), dotted=True, seed=4)
    draw_hand_text(s2, "画画", (1040, 70), f_lab, angle=8, jitter=1.3, seed=14)
    draw_arrow(s2, (1100, 145), (box_art[2] - 40, box_art[1] + 70), dotted=True, seed=5)
    draw_hand_text(s2, "穿搭", (1688, 90), f_lab, angle=8, jitter=1.4, seed=15)
    draw_arrow(s2, (1740, 165), (box_dress[0] + 90, box_dress[1] + 50), dotted=True, seed=6)
    draw_heart(s2, 620, 180, s=16, angle=-22)
    draw_heart(s2, 1200, 980, s=13, angle=14)
    draw_star(s2, 1120, 540, s=13, angle=25)
    draw_spark(s2, 640, 980, s=11)
    draw_star(s2, 1820, 980, s=12, angle=-12)
    draw_spark(s2, 980, 980, s=9)

    p2 = ROOT / "slide2_grownup.png"
    flatten(s2).save(p2, "PNG")
    print("wrote", p2)

    # ---------- SLIDE 3 ----------
    print("slide 3...")
    s3 = new_canvas()
    st_dress3 = make_sticker(cuts["dress"], border=18, jitter_amp=5.7, seed=19, max_h=900)
    st_dress3 = rotate_keep(st_dress3, -5.8)
    st_chili = make_sticker(cuts["chili"], border=16, jitter_amp=5.4, seed=22, max_h=540, max_w=760)
    st_chili = rotate_keep(st_chili, 6.2)

    box_ch = paste(s3, st_chili, (70, 200), anchor="lt")
    box_dr = paste(s3, st_dress3, (1200, 80), anchor="lt")

    draw_hand_text(s3, "现在的你", (700, 36), f_title, angle=-2.6, jitter=2.3, seed=17)
    draw_hand_text(s3, "认真生活～闪闪发光", (60, 910), load_font(FONT_TITLE, 64), angle=-3.2, jitter=1.6, seed=18)
    f_lab3 = load_font(FONT_LABEL, 44)
    draw_hand_text(s3, "闪光日常", (760, 250), f_lab3, angle=10, jitter=1.3, seed=20)
    draw_arrow(s3, (860, 320), (box_ch[2] - 30, box_ch[1] + 70), dotted=True, seed=8)
    draw_heart(s3, 1080, 150, s=18, angle=-12)
    draw_heart(s3, 520, 140, s=13, angle=22)
    draw_star(s3, 1740, 70, s=15, angle=18)
    draw_spark(s3, 980, 620, s=13)
    draw_star(s3, 1780, 980, s=13, angle=-8)
    draw_heart(s3, 1100, 980, s=12, angle=16)

    p3 = ROOT / "slide3_now.png"
    flatten(s3).save(p3, "PNG")
    print("wrote", p3)

    print("pptx...")
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for img in (p1, p2, p3):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    out_pptx = ROOT / "preview.pptx"
    prs.save(str(out_pptx))
    print("wrote", out_pptx)
    print("DONE")


def main():
    child = os.environ.get("CHILD_IMG")
    collage = os.environ.get("COLLAGE_IMG")
    args = sys.argv[1:]
    i = 0
    while i < len(args):
        if args[i] == "--child" and i + 1 < len(args):
            child = args[i + 1]
            i += 2
        elif args[i] == "--collage" and i + 1 < len(args):
            collage = args[i + 1]
            i += 2
        else:
            i += 1
    build_slides(
        Path(child) if child else None,
        Path(collage) if collage else None,
    )


if __name__ == "__main__":
    main()
